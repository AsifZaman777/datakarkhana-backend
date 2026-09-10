import time
import re
import os
import hashlib
import threading
import queue
from datetime import datetime
from io import BytesIO
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options

SCROLL_TIMES = 10
WAIT_TIME = 2

# ── In-memory live event & frame buffer for WebSocket streaming ──
LIVE_FRAMES: dict[int, dict] = {}
_frame_lock = threading.Lock()

JOB_SUBSCRIBERS: dict[int, set] = {}
_sub_lock = threading.Lock()

JOB_RECENT_LOGS: dict[int, list[str]] = {}
_logs_lock = threading.Lock()

def add_job_log(job_id: int, message: str) -> str:
    """Record a formatted log line in memory and broadcast in real-time to active WebSocket subscribers"""
    formatted = f"[{datetime.now().strftime('%H:%M:%S')}] {message}"
    if job_id:
        jid = int(job_id)
        with _logs_lock:
            if jid not in JOB_RECENT_LOGS:
                JOB_RECENT_LOGS[jid] = []
            JOB_RECENT_LOGS[jid].append(formatted)
            if len(JOB_RECENT_LOGS[jid]) > 500:
                JOB_RECENT_LOGS[jid] = JOB_RECENT_LOGS[jid][-300:]

        publish_scraper_event(jid, {"type": "log", "message": formatted})
    return formatted

def get_job_recent_logs(job_id: int) -> list[str]:
    """Retrieve in-memory logs for a job without touching the database"""
    if not job_id:
        return []
    with _logs_lock:
        return list(JOB_RECENT_LOGS.get(int(job_id), []))

def subscribe_scraper_job(job_id: int) -> queue.Queue:
    """Subscribe a WebSocket worker to real-time events for a specific scraper job"""
    q = queue.Queue(maxsize=150)
    jid = int(job_id)
    with _sub_lock:
        if jid not in JOB_SUBSCRIBERS:
            JOB_SUBSCRIBERS[jid] = set()
        JOB_SUBSCRIBERS[jid].add(q)
    return q

def unsubscribe_scraper_job(job_id: int, q: queue.Queue):
    """Unsubscribe a WebSocket worker"""
    if not job_id:
        return
    jid = int(job_id)
    with _sub_lock:
        if jid in JOB_SUBSCRIBERS:
            JOB_SUBSCRIBERS[jid].discard(q)
            if not JOB_SUBSCRIBERS[jid]:
                JOB_SUBSCRIBERS.pop(jid, None)

def publish_scraper_event(job_id: int, event: dict):
    """Push an event to all live subscribers of a job without touching the database"""
    if not job_id:
        return
    jid = int(job_id)
    with _sub_lock:
        subscribers = list(JOB_SUBSCRIBERS.get(jid, []))
    for q in subscribers:
        try:
            q.put_nowait(event)
        except queue.Full:
            try:
                q.get_nowait()
                q.put_nowait(event)
            except Exception:
                pass

def push_scraper_frame(driver, job_id):
    """Capture a compressed JPEG frame and store it in the in-memory buffer"""
    if not job_id:
        return
    jid = int(job_id)
    try:
        # Get screenshot as PNG bytes from Selenium
        png_bytes = driver.get_screenshot_as_png()

        # Compress to JPEG using Pillow for smaller frames (~20-50KB vs ~500KB PNG)
        from PIL import Image as PILImage
        img = PILImage.open(BytesIO(png_bytes))

        # Resize to max 800px width to keep frames lightweight
        max_width = 800
        if img.width > max_width:
            ratio = max_width / img.width
            img = img.resize((max_width, int(img.height * ratio)), PILImage.LANCZOS)

        buffer = BytesIO()
        img.save(buffer, format="JPEG", quality=40, optimize=True)
        jpeg_bytes = buffer.getvalue()

        import base64
        b64_data = base64.b64encode(jpeg_bytes).decode("utf-8")
        frame_hash = hashlib.md5(jpeg_bytes).hexdigest()

        with _frame_lock:
            LIVE_FRAMES[jid] = {"data": b64_data, "hash": frame_hash}
        publish_scraper_event(jid, {"type": "frame", "image": b64_data})
    except Exception:
        pass

def get_live_frame(job_id: int) -> dict | None:
    """Get the latest frame for a job (thread-safe)"""
    if not job_id:
        return None
    with _frame_lock:
        return LIVE_FRAMES.get(int(job_id))

def clear_scraper_frame(job_id: int):
    """Keep the last captured frame for user review until memory threshold is reached"""
    with _frame_lock:
        if len(LIVE_FRAMES) > 25:
            try:
                oldest = next(iter(LIVE_FRAMES))
                LIVE_FRAMES.pop(oldest, None)
            except Exception:
                pass

# Backward-compatible alias — all existing call sites use this name
save_scraper_screenshot = push_scraper_frame

# ── Manual Job Interruption / Stop Registry ──
STOPPED_JOBS: set[int] = set()
_stop_lock = threading.Lock()

def stop_scraper_job(job_id: int):
    """Signal a job to stop scraping immediately"""
    with _stop_lock:
        STOPPED_JOBS.add(job_id)

def is_job_stopped(job_id: int) -> bool:
    """Check if job has received a stop signal"""
    if not job_id:
        return False
    with _stop_lock:
        return job_id in STOPPED_JOBS

def clear_job_stop(job_id: int):
    """Clean up stop registry entry"""
    with _stop_lock:
        STOPPED_JOBS.discard(job_id)



def setup_driver(headless=True):
    """Setup Chrome in background mode with Selenium and anti-interception protection"""
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-popup-blocking")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-infobars")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    options.add_argument("--window-size=1400,900")
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
    driver = webdriver.Chrome(options=options)
    driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")

    # Prevent reload interception & beforeunload popups via CDP
    try:
        driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {
            "source": """
                window.location.reload = function() { console.log('[Scraper Protection] Page reload suppressed.'); };
                window.addEventListener('beforeunload', function(e) {
                    e.stopImmediatePropagation();
                }, true);
            """
        })
    except Exception:
        pass

    return driver

def extract_phone(text):
    """Extract Bangladesh phone number from page source or elements"""
    patterns = [
        r'(\+?88001[3-9]\d{8})',
        r'(\+?8801[3-9]\d{8})',
        r'(8801[3-9]\d{8})',
        r'(01[3-9]\d{8})',
        r'(1[3-9]\d{8})',
    ]
    for pattern in patterns:
        match = re.search(pattern, text.replace('-', '').replace(' ', ''))
        if match:
            num = match.group(1)
            if num.startswith('+88001'):
                num = '+8801' + num[6:]
            elif num.startswith('88001'):
                num = '+8801' + num[5:]
            elif num.startswith('01'):
                num = '+88' + num
            elif num.startswith('1') and len(num) == 10:
                num = '+880' + num
            elif not num.startswith('+'):
                num = '+' + num
            return num
    return ''

def scroll_results(driver, scroll_times=8, log_cb=print, job_id=None):
    """Scroll Google Maps sidebar panel with immediate live visual feedback"""
    try:
        panel = driver.find_element(By.CSS_SELECTOR, 'div[role="feed"]')
        for i in range(scroll_times):
            driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight", panel)
            log_cb(f"📜 [SCROLL {i+1}/{scroll_times}] Feed scroll performed. Loading more places...")
            if job_id:
                push_scraper_frame(driver, job_id)
            time.sleep(1.0)
    except Exception as e:
        log_cb(f"⚠️ [SCROLL] Sidebar scroll skipped: {e}")

def scrape_query(driver, query, log_cb=print, job_id=None):
    """Scrape a search query from Google Maps with live step-by-step debugger tracking"""
    results = []
    url = f"https://www.google.com/maps/search/{query.replace(' ', '+')}"
    log_cb(f"🌐 [NAVIGATE] Loading Google Maps URL: {url}")

    driver.get(url)
    if job_id:
        push_scraper_frame(driver, job_id)
    time.sleep(2.5)
    log_cb(f"📍 [RENDERED] Google Maps page loaded for query: '{query}'")
    if job_id:
        push_scraper_frame(driver, job_id)

    # Scroll sidebar with live frame updates
    scroll_results(driver, 8, log_cb, job_id=job_id)
    if job_id:
        push_scraper_frame(driver, job_id)

    listings = driver.find_elements(By.CSS_SELECTOR, 'a[href*="/maps/place/"]')
    log_cb(f"🎯 [DISCOVERY] Found {len(listings)} place links on page. Beginning item extraction...")

    seen_names = set()

    for i, listing in enumerate(listings):
        if is_job_stopped(job_id):
            log_cb("⏹️ [HALT] Interruption / Stop signal received. Halting current scrape query...")
            break

        try:
            name = listing.get_attribute("aria-label") or ""
            href = listing.get_attribute("href") or ""

            if not name or name in seen_names:
                continue
            seen_names.add(name)

            log_cb(f"🔍 [INSPECT #{len(results)+1}] Clicking '{name}'...")
            # Click to load details
            driver.execute_script("arguments[0].click();", listing)
            if job_id:
                push_scraper_frame(driver, job_id)
            time.sleep(1.2)
            if job_id:
                push_scraper_frame(driver, job_id)

            phone = ''
            address = ''
            website = ''
            rating = ''
            category = ''

            # Phone Extraction
            try:
                phone_els = driver.find_elements(
                    By.CSS_SELECTOR,
                    'button[data-item-id*="phone"], button[aria-label*="phone"], [data-tooltip*="Copy phone"]'
                )
                for el in phone_els:
                    text = el.get_attribute("aria-label") or el.text
                    phone = extract_phone(text)
                    if phone:
                        break
            except Exception:
                pass

            # Address Extraction
            try:
                addr_els = driver.find_elements(
                    By.CSS_SELECTOR, 'button[data-item-id*="address"], [data-tooltip*="Copy address"]'
                )
                for el in addr_els:
                    txt = el.get_attribute("aria-label") or el.text
                    if txt and len(txt) > 5:
                        address = txt.replace("Address: ", "").strip()
                        break
            except Exception:
                pass

            # Website Extraction
            try:
                web_els = driver.find_elements(By.CSS_SELECTOR, 'a[data-item-id*="authority"]')
                for el in web_els:
                    w = el.get_attribute("href") or ""
                    if w and "google" not in w:
                        website = w
                        break
            except Exception:
                pass

            # Rating
            try:
                rating_el = driver.find_element(By.CSS_SELECTOR, 'span[aria-label*="stars"]')
                rating = rating_el.get_attribute("aria-label") or ""
            except Exception:
                pass

            # Category
            try:
                cat_el = driver.find_element(By.CSS_SELECTOR, 'button[jsaction*="category"]')
                category = cat_el.text.strip()
            except Exception:
                pass

            entry = {
                "Name": name,
                "Phone": phone,
                "Address": address,
                "Website": website,
                "Rating": rating,
                "Category": category,
                "Maps URL": href,
                "Query": query,
            }
            results.append(entry)

            phone_badge = f"📞 {phone}" if phone else "⚠️ No phone"
            addr_snippet = f"| 📍 {address[:30]}" if address else ""
            log_cb(f"✅ [SAVED #{len(results)}] '{name}' | {phone_badge} {addr_snippet}")

            publish_scraper_event(job_id, {
                "type": "progress",
                "count": len(results),
                "action": f"Extracted #{len(results)}: {name}",
                "item": {"name": name, "phone": phone, "category": category, "address": address}
            })
            if job_id:
                push_scraper_frame(driver, job_id)

        except Exception as e:
            log_cb(f"⚠️ [WARN] Error parsing listing #{i+1}: {e}")
            continue

    return results

def save_to_excel(all_results, filename):
    """Save results to format-aligned Excel spreadsheet and return deduplicated row count"""
    if not all_results:
        return 0
    df = pd.DataFrame(all_results)
    if "Phone" in df.columns:
        df["Phone"] = df["Phone"].astype(str).str.replace(r'\.0$', '', regex=True).replace({'nan': '', 'None': ''})
    df = df.drop_duplicates(subset=["Name", "Phone"])

    # Sort with phone numbers first
    df["has_phone"] = df["Phone"].apply(lambda x: 0 if x else 1)
    df = df.sort_values(["has_phone", "Name"]).drop(columns=["has_phone"])

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Coaching Centers")
        ws = writer.sheets["Coaching Centers"]

        col_widths = {
            "A": 35, "B": 18, "C": 40, "D": 30, "E": 10, "F": 20, "G": 50, "H": 35
        }
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        from openpyxl.styles import PatternFill, Font, Alignment
        header_fill = PatternFill(start_color="0D1B3E", end_color="0D1B3E", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        green_fill = PatternFill(start_color="E1F5EE", end_color="E1F5EE", fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            if row[1].value:  # has phone
                for cell in row:
                    cell.fill = green_fill

    return len(df)
